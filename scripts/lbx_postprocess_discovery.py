#!/usr/bin/env python3
"""Post-process LBX discovery outputs into compact indexes and extract possible store rows."""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

ROOT = Path("fast_output")
OUT = Path("post_output")
OUT.mkdir(exist_ok=True)


def load_json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def compact(value: Any, limit: int = 500) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text[:limit]


def write_tsv(path: Path, rows: Iterable[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, delimiter="\t", extrasaction="ignore")
        w.writeheader()
        for row in rows:
            w.writerow({k: compact(row.get(k, ""), 2000) for k in fields})


def summarize_candidates() -> list[dict[str, Any]]:
    candidates = load_json(ROOT / "bidding_candidates.json", [])
    rows: list[dict[str, Any]] = []
    for index, item in enumerate(candidates, start=1):
        inspection = item.get("inspection") or {}
        sheets = inspection.get("sheets") or []
        sheet_summary = []
        first_samples = []
        all_hits: set[str] = set()
        max_rows = 0
        max_cols = 0
        addr_rows = 0
        for sheet in sheets:
            max_rows = max(max_rows, int(sheet.get("max_row") or 0))
            max_cols = max(max_cols, int(sheet.get("max_col") or 0))
            addr_rows = max(addr_rows, int(sheet.get("address_like_rows") or 0))
            all_hits.update(sheet.get("keyword_hits") or [])
            sheet_summary.append(
                f"{sheet.get('name')}[{sheet.get('max_row')}x{sheet.get('max_col')};地址样式{sheet.get('address_like_rows')}]"
            )
            samples = sheet.get("samples") or []
            if samples and len(first_samples) < 5:
                first_samples.extend([" | ".join(map(str, row[:20])) for row in samples[: max(0, 5-len(first_samples))]])
        rows.append({
            "rank": index,
            "score": item.get("candidate_score"),
            "max_rows": max_rows,
            "max_cols": max_cols,
            "address_like_rows": addr_rows,
            "label": item.get("label"),
            "notice_title": item.get("notice_title"),
            "notice_url": item.get("notice_url"),
            "attachment_url": item.get("url"),
            "saved_path": item.get("saved_path"),
            "keywords": ",".join(sorted(all_hits)),
            "sheets": "; ".join(sheet_summary),
            "samples": " || ".join(first_samples),
        })
    rows.sort(key=lambda r: (-(int(r.get("score") or 0)), -(int(r.get("max_rows") or 0)), r.get("label") or ""))
    write_tsv(OUT / "candidate_index.tsv", rows, [
        "rank", "score", "max_rows", "max_cols", "address_like_rows", "label", "notice_title",
        "keywords", "sheets", "samples", "saved_path", "notice_url", "attachment_url",
    ])
    return rows


def summarize_web_apps() -> dict[str, Any]:
    data = load_json(ROOT / "web_app_discovery.json", {"pages": [], "scripts": []})
    page_rows: list[dict[str, Any]] = []
    for page in data.get("pages", []):
        page_rows.append({
            "status": page.get("status"), "url": page.get("url"), "final_url": page.get("final_url"),
            "title": page.get("title"), "content_type": page.get("content_type"), "length": page.get("length"),
            "scripts": " | ".join(page.get("scripts") or []),
            "links": " | ".join(page.get("links") or []),
            "contexts": " || ".join(page.get("contexts") or []),
            "preview": page.get("preview"),
        })
    write_tsv(OUT / "web_pages.tsv", page_rows, [
        "status", "url", "final_url", "title", "content_type", "length", "scripts", "links", "contexts", "preview"
    ])

    endpoint_rows: list[dict[str, Any]] = []
    for script in data.get("scripts", []):
        for endpoint in script.get("candidate_endpoints") or []:
            endpoint_rows.append({
                "script_url": script.get("url"), "endpoint": endpoint,
                "contexts": " || ".join(script.get("contexts") or []),
            })
        if script.get("contexts") and not script.get("candidate_endpoints"):
            endpoint_rows.append({
                "script_url": script.get("url"), "endpoint": "",
                "contexts": " || ".join(script.get("contexts") or []),
            })
    # Scan saved JS source more aggressively for quoted URL/path strings.
    path_re = re.compile(r"['\"]([^'\"]{3,500})['\"]")
    interest_re = re.compile(
        r"(store|shop|branch|outlet|location|nearby|distance|longitude|latitude|lng|lat|poi|门店|药店|药房|地址)", re.I
    )
    for js_path in ROOT.glob("js_*.txt"):
        text = js_path.read_text(encoding="utf-8", errors="ignore")
        seen: set[str] = set()
        for match in path_re.finditer(text):
            value = match.group(1)
            if not interest_re.search(value):
                continue
            if value in seen:
                continue
            seen.add(value)
            start = max(0, match.start() - 250)
            end = min(len(text), match.end() + 250)
            endpoint_rows.append({
                "script_url": js_path.name,
                "endpoint": value,
                "contexts": compact(text[start:end], 1000),
            })
            if len(seen) >= 1000:
                break
    # Deduplicate.
    unique: dict[tuple[str, str], dict[str, Any]] = {}
    for row in endpoint_rows:
        unique[(row.get("script_url") or "", row.get("endpoint") or "")] = row
    endpoint_rows = list(unique.values())
    endpoint_rows.sort(key=lambda r: (0 if "/" in (r.get("endpoint") or "") else 1, r.get("endpoint") or ""))
    write_tsv(OUT / "endpoint_hits.tsv", endpoint_rows, ["endpoint", "script_url", "contexts"])

    # Extract likely API host/base URL strings.
    host_rows: list[dict[str, Any]] = []
    host_re = re.compile(r"https?://[A-Za-z0-9._:-]+(?:/[A-Za-z0-9_./?=&%:-]*)?")
    for js_path in ROOT.glob("js_*.txt"):
        text = js_path.read_text(encoding="utf-8", errors="ignore")
        for url in sorted(set(host_re.findall(text))):
            if any(k in url.lower() for k in ["lbx", "api", "mall", "store", "shop"]):
                host_rows.append({"source": js_path.name, "url": url})
    write_tsv(OUT / "host_hits.tsv", host_rows, ["url", "source"])
    return {"page_count": len(page_rows), "endpoint_count": len(endpoint_rows), "host_count": len(host_rows)}


def inspect_mobile_map() -> dict[str, Any]:
    path = ROOT / "map_baidu_mobile.txt"
    if not path.exists():
        return {"exists": False}
    text = path.read_text(encoding="utf-8", errors="ignore")
    terms = ["老百姓大药房", "老百姓健康药房", "content", "poi", "uid", "addr", "address", "wd", "qt=s"]
    contexts: list[dict[str, Any]] = []
    for term in terms:
        start = 0
        count = 0
        while count < 50:
            idx = text.find(term, start)
            if idx < 0:
                break
            contexts.append({"term": term, "offset": idx, "context": compact(text[max(0, idx-500): idx+len(term)+1000], 1600)})
            count += 1
            start = idx + len(term)
    write_tsv(OUT / "mobile_map_contexts.tsv", contexts, ["term", "offset", "context"])
    return {"exists": True, "length": len(text), "term_counts": {t: text.count(t) for t in terms}, "context_count": len(contexts)}


def looks_like_header(values: list[str]) -> bool:
    text = "|".join(values)
    groups = [
        any(x in text for x in ["门店", "店名", "药店", "药房", "网点"]),
        any(x in text for x in ["地址", "省", "市", "区县", "区域"]),
    ]
    return all(groups)


def extract_store_like_sheets() -> dict[str, Any]:
    attachment_dir = ROOT / "candidate_attachments"
    inventory: list[dict[str, Any]] = []
    extracted: list[dict[str, Any]] = []
    for path in sorted(attachment_dir.glob("*.xlsx")):
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            inventory.append({"file": path.name, "error": repr(exc)})
            continue
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = ["" if v is None else compact(v, 1000) for v in row]
                if any(vals):
                    rows.append(vals)
                if r_idx >= 25000:
                    break
            header_idx = None
            for i, vals in enumerate(rows[:50]):
                if looks_like_header(vals):
                    header_idx = i
                    break
            inv = {
                "file": path.name, "sheet": ws.title, "max_row": ws.max_row, "max_col": ws.max_column,
                "nonempty_rows": len(rows), "header_row_index": (header_idx + 1 if header_idx is not None else ""),
                "sample": " || ".join(" | ".join(r[:25]) for r in rows[:6]),
            }
            inventory.append(inv)
            if header_idx is not None:
                header = rows[header_idx]
                for data_index, vals in enumerate(rows[header_idx+1:], start=header_idx+2):
                    if not any(vals):
                        continue
                    row_dict = {header[i] or f"列{i+1}": vals[i] if i < len(vals) else "" for i in range(max(len(header), len(vals)))}
                    row_dict = {k: v for k, v in row_dict.items() if k or v}
                    extracted.append({
                        "source_file": path.name,
                        "source_sheet": ws.title,
                        "source_row": data_index,
                        "row_json": json.dumps(row_dict, ensure_ascii=False),
                    })
        wb.close()
    write_tsv(OUT / "attachment_sheet_inventory.tsv", inventory, [
        "file", "sheet", "max_row", "max_col", "nonempty_rows", "header_row_index", "sample", "error"
    ])
    write_tsv(OUT / "extracted_store_like_rows.tsv", extracted, [
        "source_file", "source_sheet", "source_row", "row_json"
    ])
    return {"sheet_count": len(inventory), "extracted_row_count": len(extracted)}


def main() -> None:
    candidates = summarize_candidates()
    web = summarize_web_apps()
    mobile = inspect_mobile_map()
    sheets = extract_store_like_sheets()
    summary = {
        "candidate_count": len(candidates),
        "top_candidates": candidates[:30],
        "web": web,
        "mobile_map": mobile,
        "attachments": sheets,
    }
    (OUT / "post_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2)[:30000])


if __name__ == "__main__":
    main()
